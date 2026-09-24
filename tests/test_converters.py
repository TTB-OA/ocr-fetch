"""Conversion checks across the bundled fixture documents."""

from pathlib import Path

import pytest
from conftest import fixture_paths, requires_poppler, requires_tesseract

from ocr_fetch import (
    can_convert,
    convert_file_to_markdown,
    get_parse_method_name,
    is_conversion_error,
    looks_like_text,
    supported_extensions,
)
from ocr_fetch.errors import DEPENDENCY_ERROR_PREFIX


def _assert_converts(path: Path) -> str:
    content, method = convert_file_to_markdown(str(path))
    if str(content).strip().startswith(DEPENDENCY_ERROR_PREFIX):
        pytest.skip(f"Missing dependency for {path.name}: {content}")
    assert not is_conversion_error(content), f"{path.name} -> {method}: {content}"
    assert content.strip(), f"{path.name} -> {method}: empty output"
    return method


@pytest.mark.parametrize("path", fixture_paths(".txt", ".csv", ".htm", ".html", ".xml"), ids=lambda p: p.name)
def test_text_like_fixtures(path: Path):
    _assert_converts(path)


@pytest.mark.parametrize("path", fixture_paths(".xls", ".xlsx"), ids=lambda p: p.name)
def test_spreadsheet_fixtures(path: Path):
    _assert_converts(path)


@pytest.mark.parametrize("path", fixture_paths(".doc", ".docx", ".pptx"), ids=lambda p: p.name)
def test_word_fixtures(path: Path):
    _assert_converts(path)


@requires_tesseract
@requires_poppler
@pytest.mark.parametrize("path", fixture_paths(".pdf"), ids=lambda p: p.name)
def test_pdf_fixtures(path: Path):
    _assert_converts(path)


@requires_tesseract
@pytest.mark.parametrize("path", fixture_paths(".jpg", ".jpeg", ".png", ".bmp"), ids=lambda p: p.name)
def test_image_fixtures(path: Path):
    _assert_converts(path)


def test_pptx_converts(tmp_path: Path):
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Quarterly wine labeling update"
    path = tmp_path / "deck.pptx"
    deck.save(str(path))
    content, method = convert_file_to_markdown(str(path))
    assert method == "markitdown_pptx"
    assert "Quarterly wine labeling update" in content


def test_unsupported_extension_falls_back_to_text(tmp_path: Path):
    odd = tmp_path / "notes.unknownext"
    odd.write_text("hello world", encoding="utf-8")
    content, method = convert_file_to_markdown(str(odd))
    assert method == "plaintext"
    assert "hello world" in content


def test_unregistered_binary_returns_error_marker(tmp_path: Path):
    blob = tmp_path / "payload.bin"
    blob.write_bytes(b"PK\x03\x04\x14\x00" + bytes(range(256)) * 8)
    content, method = convert_file_to_markdown(str(blob), content_type="application/octet-stream")
    assert is_conversion_error(content)
    assert "Unsupported binary" in content
    assert method == "unknown"


def test_unknown_extension_missing_file_returns_error_marker(tmp_path: Path):
    content, method = convert_file_to_markdown(str(tmp_path / "ghost.unknownext"))
    assert is_conversion_error(content)
    assert method == "unknown"


def test_looks_like_text(tmp_path: Path):
    text = tmp_path / "a.dat"
    text.write_text("plain\ttext\r\nwith newlines\n", encoding="utf-8")
    assert looks_like_text(str(text))
    binary = tmp_path / "b.dat"
    binary.write_bytes(b"abc\x00def")
    assert not looks_like_text(str(binary))
    empty = tmp_path / "c.dat"
    empty.write_bytes(b"")
    assert looks_like_text(str(empty))


def test_missing_file_returns_error_marker(tmp_path: Path):
    content, _ = convert_file_to_markdown(str(tmp_path / "does_not_exist.txt"))
    assert is_conversion_error(content)


@pytest.mark.parametrize(
    ("file_ext", "content_type", "expected"),
    [
        (".pdf", None, "pypdf_text"),
        (".HTM", None, "html2text"),
        ("", "text/csv; charset=utf-8", "pandas_markdown"),
        (".zip", None, "zip_archive"),
        ("", "application/zip", "zip_archive"),
        (".zzz", None, "unknown"),
    ],
)
def test_get_parse_method_name(file_ext: str, content_type: str | None, expected: str):
    assert get_parse_method_name(file_ext, content_type) == expected


@pytest.mark.parametrize(
    ("file_ext", "content_type", "expected"),
    [
        (".pdf", None, True),
        ("PDF", None, True),
        (".md", None, True),
        (".xsd", None, True),
        ("", "application/zip; charset=binary", True),
        (".zzz", "application/octet-stream", False),
        (None, None, False),
    ],
)
def test_can_convert(file_ext: str | None, content_type: str | None, expected: bool):
    assert can_convert(file_ext, content_type) is expected


def test_supported_extensions_matches_registry():
    exts = supported_extensions()
    assert exts == tuple(sorted(exts))
    assert {".pdf", ".zip", ".txt"} <= set(exts)
    assert all(e.startswith(".") and e == e.lower() for e in exts)
